# -*- coding: utf-8 -*-
"""LAS LINEAS SIN TIPO (Monica, 21-oct-2026): para Negocio > Trabajo contratado, que
agrupa lo contratado por lo que es cada linea (conceptos_hoja.tipo_proyecto_id). Las
lineas de una hoja firmada que llevan dinero (importe o %) y no dicen que son salen
en «Sin tipo apuntado». Este Excel las lista para que ella les ponga el tipo a mano
(columna «Tipo», con desplegable de tipos_proyecto). SOLO LEE: no escribe nada.

  python scripts/negocio/lineas_sin_tipo.py   ->  docs/negocio/lineas_sin_tipo.xlsx
"""
import os, sys, collections
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
from produccion import arrancar  # noqa: E402
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

b = arrancar()
tipos = b.leer('tipos_proyecto?select=nombre&activo=eq.true&elegible=eq.true&order=orden.asc')
hojas = b.leer('hojas_encargo?select=id,numero_hoja,fecha_firma,version_firmada_id,oportunidad:oportunidad_id(nombre,estado)'
               '&estado=eq.devuelta_firmada&version_firmada_id=not.is.null')
hojas = [h for h in hojas if h['oportunidad'] and h['oportunidad']['estado'] != 'anulada']
por_version = {h['version_firmada_id']: h for h in hojas}
lineas = []
vs = list(por_version)
for i in range(0, len(vs), 150):
    lineas += b.leer('conceptos_hoja?select=id,version_hoja_id,descripcion,importe,porcentaje,bloque:bloque_id(nombre)'
                     '&tipo_proyecto_id=is.null&incluido_en_concepto_id=is.null&version_hoja_id=in.(%s)' % ','.join(vs[i:i + 150]))
filas = []
for l in lineas:
    imp, pc = float(l['importe'] or 0), float(l['porcentaje'] or 0)
    if imp <= 0 and pc <= 0:
        continue  # lo incluido (0 € sin tipo): no cuenta como contratado
    h = por_version[l['version_hoja_id']]
    filas.append([(h['fecha_firma'] or '')[:4], (h['fecha_firma'] or '')[:10], h['numero_hoja'] or '', h['oportunidad']['nombre'] or '',
                  (l['bloque'] or {}).get('nombre', ''), (l['descripcion'] or '').strip(), imp or None, pc or None, '', l['id']])
filas.sort(key=lambda f: (f[1], f[2]), reverse=True)

wb = Workbook()
ws = wb.active
ws.title = 'Sin tipo'
cab = ['Año', 'Firmada', 'Hoja', 'Oportunidad', 'Bloque', 'Lo que dice la línea', 'Importe', '%', 'Tipo (ponlo tú)', 'id (no tocar)']
ws.append(cab)
for f in filas:
    ws.append(f)
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='2E3236')
for c in ws['I'][1:]:
    c.fill = PatternFill('solid', fgColor='FFF4C6')
for col, ancho in zip('ABCDEFGHIJ', [6, 11, 12, 34, 34, 60, 11, 6, 28, 38]):
    ws.column_dimensions[col].width = ancho
for c in ws['F'][1:]:
    c.alignment = Alignment(wrap_text=True, vertical='top')
for c in ws['G'][1:]:
    c.number_format = '#,##0.00 €'
ws.freeze_panes = 'A2'
ws.auto_filter.ref = 'A1:J%d' % (len(filas) + 1)

lista = wb.create_sheet('Tipos')
for t in tipos:
    lista.append([t['nombre']])
dv = DataValidation(type='list', formula1='=Tipos!$A$1:$A$%d' % len(tipos), allow_blank=True)
ws.add_data_validation(dv)
dv.add('I2:I%d' % (len(filas) + 1))

res = wb.create_sheet('Resumen', 0)
res.append(['Año', 'Bloque', 'Líneas', 'Importe'])
cuenta = collections.defaultdict(lambda: [0, 0.0])
for f in filas:
    k = (f[0], f[4]); cuenta[k][0] += 1; cuenta[k][1] += f[6] or 0
for (a, bl), (n, imp) in sorted(cuenta.items(), key=lambda x: (x[0][0], -x[1][1]), reverse=False):
    res.append([a, bl, n, imp])
for c in res[1]:
    c.font = Font(bold=True)
res.column_dimensions['B'].width = 48
for c in res['D'][1:]:
    c.number_format = '#,##0 €'

salida = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..', 'docs', 'negocio', 'lineas_sin_tipo.xlsx')
os.makedirs(os.path.dirname(salida), exist_ok=True)
wb.save(salida)
print(len(filas), 'lineas;', sum(1 for f in filas if f[0] == '2026'), 'de 2026 ->', os.path.normpath(salida))

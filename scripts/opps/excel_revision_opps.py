# -*- coding: utf-8 -*-
"""Lo que Mónica tiene que mirar de la lista de opps (20-oct-2026), en un Excel: docs/opps/revisar_opps_20oct.xlsx
  1. Portales de PARCELAS DISTINTAS (o nombre con fase/bloque/portal): ¿son de la opp?
  2. Nombre a mano con letra o «bis» que Catastro no tiene (se quedan con su nombre).
  3. Comunidades sin ninguna opp (cáscaras: no existe comunidad sin opp).
  4. Opps vivas sin comunidad.
  5. Opps sin nombre ni portal (no se les puede poner nombre de Catastro).
Solo lee; no escribe nada en la base.

  python scripts/opps/excel_revision_opps.py
"""
import sys, os, re, json, collections, unicodedata
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
from produccion import arrancar  # noqa: E402
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

RAIZ = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..')
b = arrancar()
ops = {o['id']: o for o in b.leer('oportunidades?select=id,codigo,nombre,estado,comunidad_id&order=codigo')}
vivas = {k: o for k, o in ops.items() if o['estado'] != 'anulada'}
acc = {a['id']: a for a in b.leer('accesos?select=id,ficha_catastro_portal_id,figura_legal_propietaria_id&order=id')}
rel = collections.defaultdict(list)
for r in b.leer('relacion_oportunidad_accesos?select=opp_id,acceso_id&hasta=is.null&order=opp_id'): rel[r['opp_id']].append(acc[r['acceso_id']])
hojas = collections.Counter(h['oportunidad_id'] for h in b.leer('hojas_encargo?select=oportunidad_id&order=id') if h['oportunidad_id'])
facturas = collections.Counter(f['oportunidad_id'] for f in b.leer('facturas?select=oportunidad_id&tipo=eq.factura&order=id') if f['oportunidad_id'])
por_codigo = {o['codigo']: o for o in vivas.values()}

wb = Workbook()
def hoja(titulo, cabecera, filas, anchos):
    ws = wb.create_sheet(titulo[:31])
    ws.append(cabecera)
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='3D3D3D'); c.alignment = Alignment(wrap_text=True)
    for f in filas: ws.append(f)
    for i, w in enumerate(anchos): ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = 'A2'
    return ws

rev = json.load(open(os.path.join(RAIZ, 'docs', 'opps', 'revisar_portales_parcelas.json'), encoding='utf-8'))
hoja('1 Portales a revisar', ['Código', 'Estado', 'Nombre que tiene', 'Lo que dicen sus portales (Catastro)', 'Por qué', 'Parcelas', 'Hojas', 'Facturas', 'Tu decisión'],
     [[r['codigo'], r['estado'], r['antes'], r['catastro'], r['motivo'], r['parcelas'], hojas.get(por_codigo[r['codigo']]['id'], 0),
       facturas.get(por_codigo[r['codigo']]['id'], 0), ''] for r in rev], [14, 9, 40, 48, 40, 34, 7, 8, 30])

def norm(s): return ' '.join(re.sub(r'[^a-z0-9 ]', ' ', ''.join(c for c in unicodedata.normalize('NFD', (s or '').lower()) if unicodedata.category(c) != 'Mn')).split())
letra = [o for o in vivas.values() if o['nombre'] and o['nombre'] == o['nombre'].upper() and re.search(r'\b\d+\s*(bis|[a-z])\b(?!\s*-?\s*\d)', norm(o['nombre']))
         and rel.get(o['id'])]
hoja('2 Letra o bis (no se tocan)', ['Código', 'Estado', 'Nombre (se queda así)'], [[o['codigo'], o['estado'], o['nombre']] for o in letra], [14, 9, 50])

coms = b.leer('comunidades?select=id,nombre,creado_en,cif_comunidad&order=nombre')
con = {o['comunidad_id'] for o in vivas.values() if o['comunidad_id']}
nacc = collections.Counter(a['figura_legal_propietaria_id'] for a in acc.values())
hoja('3 Comunidades sin opp', ['Comunidad', 'Alta', 'CIF', 'Portales', 'Hojas', 'Tu decisión'],
     [[c['nombre'], c['creado_en'][:10], c['cif_comunidad'] or '', nacc.get(c['id'], 0), len(b.leer('hojas_encargo?select=id&comunidad_id=eq.' + c['id'], False)), '']
      for c in coms if c['id'] not in con], [48, 11, 12, 9, 7, 30])

hoja('4 Opps sin comunidad', ['Código', 'Nombre', 'Estado', 'Portales', 'Hojas', 'Facturas', 'Tu decisión'],
     [[o['codigo'], o['nombre'] or '(sin nombre)', o['estado'], len(rel.get(o['id'], [])), hojas.get(o['id'], 0), facturas.get(o['id'], 0), '']
      for o in vivas.values() if not o['comunidad_id']], [14, 44, 9, 9, 7, 8, 30])

hoja('5 Sin nombre ni portal', ['Código', 'Estado', 'Hojas', 'Facturas'],
     [[o['codigo'], o['estado'], hojas.get(o['id'], 0), facturas.get(o['id'], 0)] for o in vivas.values() if not o['nombre'] and not rel.get(o['id'])], [14, 9, 7, 8])
wb.remove(wb['Sheet'])
destino = os.path.join(RAIZ, 'docs', 'opps', 'revisar_opps_20oct.xlsx')
wb.save(destino)
for ws in wb.worksheets: print('%-30s %d filas' % (ws.title, ws.max_row - 1))
print(destino)

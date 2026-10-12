# -*- coding: utf-8 -*-
"""POR QUE NO CASAN las lineas de factura que siguen sin hito (Monica, 12-oct-2026): "no se casa porque no esta la
hoja o porque no se han generado bien las lineas facturables?". Solo las sueltas: las que no estan en
factura_linea_hitos (fuera las de Ecobalance, abonos y lo que no es de una hoja).

Para cada una dice la CAUSA:
  - FALTA LA HOJA: no se encuentra la comunidad, la comunidad no tiene hoja firmada, o la factura cita una hoja
    (por su fecha) que no esta.
  - LINEAS MAL GENERADAS: la hoja esta, el concepto esta en la hoja (conceptos de la version) pero no hay su linea
    facturable; o la hoja no tiene ninguna linea.
  - LA HOJA NO LO INCLUYE: la hoja esta pero ese concepto no aparece en ella (otra hoja, o un extra sin hoja).
  - HITOS MAL GENERADOS: la linea esta, pero ninguno de sus hitos es de ese plazo ni de ese importe.
  - CASA AHORA / SUBHITO / VARIAS OPCIONES: se escribiria (con su OK).

SOLO LEE. Deja un Excel en Descargas.

  python scripts/facturacion/por_que_no_casan.py <carpeta con extraido_*.jsonl>
"""
import sys, os, re, json, glob, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

sys.path.insert(0, os.path.dirname(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
from produccion import arrancar  # noqa: E402
import casar_facturas_hitos as C  # noqa: E402

SALIDA = os.path.expanduser(r'~\Downloads\por_que_no_casan.xlsx')


def limpia(s):
    # O'DONNELL -> ODONNELL (el k() del 10-oct lo partia en dos y no casaba con la comunidad)
    return re.sub(r"['’´`]", '', str(s or ''))


def familias(desc):
    """Lo que cubre una linea o concepto de la hoja. Una linea puede llevar VARIAS cosas juntas ("Proyecto PORTAL
    ACCESIBILIDAD + CSS", "Proyecto, Tramitacion de Licencia, DF, CSS"): el familia_de_linea del 10-oct le daba una
    sola y la daba por ausente."""
    d = C.k(desc)
    out = set()
    if re.search(r'\bCSS\b|\bCSYS\b|SEGURIDAD Y SALUD|RECURSO PREVENTIVO', d): out.add('css')   # (12-oct) CSYS en los pedidos de FAIN
    if re.search(r'SUBVENCI|DOCUMENTACION TECNICA', d): out.add('subvencion')
    if re.search(r'DIRECCION (DE OBRA|FACULTATIVA)|\bDF\b|\bDO\b', d): out.add('df')
    if re.search(r'\bIEE\b|\bITE\b|EVALUACION DEL EDIFICIO|LIBRO DEL EDIFICIO|CERTIFICADO DE EFICIENCIA|\bCEE\b', d): out.add('iee')
    if re.search(r'INFORME|VIABILIDAD|PRESUPUESTOS|MEDICIONES', d): out.add('informe')
    if re.search(r'PROYECTO|REDACCION|ASCENSOR|PLATAFORMA|RAMPA|SATE|LICENCIA|MODELADO|TOMA DE DATOS|COTA CERO', d) or not out: out.add('proyecto')
    return out


def main():
    carpeta = sys.argv[1]
    b = arrancar()
    leidas = {}
    for f in glob.glob(os.path.join(carpeta, 'extraido_*.jsonl')):
        for ln in open(f, encoding='utf-8'):
            if ln.strip():
                x = json.loads(ln); leidas[x['id']] = x
    casadas = {x['factura_linea_id'] for x in b.leer('factura_linea_hitos?select=factura_linea_id&order=factura_linea_id,hito_cobro_id')}
    FL = [fl for fl in b.leer('factura_lineas?select=id,concepto,base,factura:factura_id!inner(tipo,empresa_emisora,anio,serie,numero,fecha,estado,pagador_nombre,pagador_nif)'
                              '&factura.serie=eq.1&factura.estado=neq.anulada&order=id') if fl['id'] not in casadas]

    def palabras(nombre):
        ws = set(C.k(limpia(nombre)).split())
        for a, z in re.findall(r'(\d+)\s*(?:AL|A)\s*(\d+)', C.k(nombre)):
            if int(a) < int(z) <= int(a) + 60: ws |= {str(n) for n in range(int(a), int(z) + 1)}
        return ws
    coms = {c['id']: (c['nombre'], palabras(c['nombre']), (c['cif_comunidad'] or '').upper()) for c in b.leer('comunidades?select=id,nombre,cif_comunidad&order=id')}
    hojas = b.leer('hojas_encargo?select=id,numero_hoja,fecha_creacion,fecha_firma,comunidad_id,estado,version_firmada_id&estado=neq.anulada&order=id')
    por_com = collections.defaultdict(list)
    for h in hojas: por_com[h['comunidad_id']].append(h)
    lineas = collections.defaultdict(list)
    for l in b.leer('lineas_facturacion?select=id,hoja_encargo_id,descripcion,importe,porcentaje,bloque:bloque_id(nombre)&order=id'):
        l['descripcion'] = l['descripcion'] or (l['bloque'] or {}).get('nombre') or ''
        lineas[l['hoja_encargo_id']].append(l)
    hitos = collections.defaultdict(list)
    for h in b.leer('hitos_cobro?select=id,linea_facturacion_id,hito,orden,porcentaje,importe,hito_padre_id&order=id'):
        if not h['hito_padre_id']: hitos[h['linea_facturacion_id']].append(h)
    conceptos = collections.defaultdict(list)
    for c in b.leer('conceptos_hoja?select=hoja_encargo_id,version_hoja_id,descripcion,importe,desglose,forma_pago,bloque:bloque_id(nombre)&order=id'):
        conceptos[c['hoja_encargo_id']].append(c)

    filas, cuenta = [], collections.Counter()
    for fl in FL:
        x = leidas.get(fl['id'], {})
        f = fl['factura']
        imp = float(fl['base'] or 0)
        fam = C.FAMILIA.get(x.get('concepto'), 'proyecto')
        if f['empresa_emisora'] == 'ecobalance' or f['tipo'] == 'abono' or fam is None or not imp:
            continue
        r = {'factura': '%s %s/%s' % (f['empresa_emisora'], f['anio'], f['numero']), 'fecha': f['fecha'], 'estado_factura': f['estado'],
             'linea': fl['id'], 'fecha_factura': f['fecha'],
             'cliente': f['pagador_nombre'], 'importe': imp, 'concepto': x.get('concepto'), 'plazo': x.get('plazo'),
             'direccion': x.get('direccion'), 'hoja_citada': x.get('fecha_hoja'), 'texto': ' '.join((fl['concepto'] or '').split())[:250],
             'comunidad': '', 'hojas': '', 'lo_que_tiene_la_hoja': '', 'hitos': ''}
        calle, num = C.calle_y_numero(limpia(x.get('direccion')))
        cand = set()
        if calle and num:
            cand = {cid for cid, (n, ws, cif) in coms.items() if all(w in ws for w in calle) and num in ws}
        nif = (f['pagador_nif'] or '').upper().replace('-', '')
        if nif:
            por_nif = {cid for cid, (n, ws, cif) in coms.items() if cif and cif.replace('-', '') == nif}
            if por_nif and (not cand or por_nif & cand): cand = (por_nif & cand) or por_nif
        r['comunidad'] = ' / '.join(coms[c][0] for c in cand)
        hs = [h for cid in cand for h in por_com.get(cid, [])]
        firmadas = [h for h in hs if h['estado'] == 'devuelta_firmada']
        if x.get('fecha_hoja'):
            misma = [h for h in hs if x['fecha_hoja'] in (h['fecha_creacion'], h['fecha_firma'])]
        else:
            misma = firmadas
        r['hojas'] = ', '.join('%s (%s, %s)' % (h['numero_hoja'], h['fecha_creacion'], 'firmada' if h['estado'] == 'devuelta_firmada' else h['estado']) for h in hs)

        def causa(c, detalle=''):
            r['causa'] = c; r['detalle'] = detalle; cuenta[c] += 1; filas.append(r)

        if not cand:
            causa('FALTA LA HOJA', 'no se encuentra la comunidad de esa dirección'); continue
        if not firmadas:
            causa('FALTA LA HOJA', 'la comunidad no tiene ninguna hoja firmada' + (' (hay enviadas sin firmar)' if hs else '')); continue
        if not misma:
            causa('FALTA LA HOJA', 'la factura cita una hoja de %s y no está' % x['fecha_hoja']); continue
        hs = misma
        r['lo_que_tiene_la_hoja'] = ' || '.join(
            '%s: %s' % (h['numero_hoja'], '; '.join('%s %s' % ((l['descripcion'] or '')[:40], l['importe'] if l['importe'] is not None else str(l['porcentaje']) + '%') for l in lineas.get(h['id'], [])) or 'SIN LÍNEAS')
            for h in hs)
        ls = [(h, l) for h in hs for l in lineas.get(h['id'], []) if fam in familias(l['descripcion'])]
        if not ls:
            sin_lineas = [h for h in hs if not lineas.get(h['id'])]
            en_conceptos = [h for h in hs if any(C.familia_de_linea(c['descripcion'] or (c['bloque'] or {}).get('nombre') or '') == fam
                                                   and c['desglose'] != 'no_aparece' for c in conceptos.get(h['id'], []))]
            if en_conceptos:
                causa('LÍNEAS MAL GENERADAS', 'el concepto está en la hoja %s pero no tiene su línea facturable' % en_conceptos[0]['numero_hoja'])
            elif sin_lineas and len(sin_lineas) == len(hs):
                causa('LÍNEAS MAL GENERADAS', 'la hoja no tiene ninguna línea facturable' + ('' if conceptos.get(sin_lineas[0]['id']) else ' (ni conceptos leídos)'))
            else:
                causa('LA HOJA NO LO INCLUYE', 'la hoja está pero no lleva %s: ¿otra hoja o un extra?' % fam)
            continue
        r['lineas_ids'] = [l['id'] for h, l in ls]; r['linea_factura_id'] = fl['id']
        tipos = C.PLAZO.get(x.get('plazo'))
        r['hitos'] = ' || '.join('%s: %s' % ((l['descripcion'] or '')[:30], ', '.join('%s %s' % (t['hito'], t['importe']) for t in hitos.get(l['id'], [])) or 'SIN HITOS') for h, l in ls)
        opciones = [(h, l, t) for h, l in ls for t in hitos.get(l['id'], []) if not tipos or t['hito'] in tipos]
        por_importe = [(h, l, t) for h, l in ls for t in hitos.get(l['id'], []) if C.cerca(t['importe'], imp)]
        if not any(hitos.get(l['id']) for h, l in ls):
            causa('HITOS MAL GENERADOS', 'la línea no tiene hitos'); continue
        # Primero el hito de ese plazo y ese importe; si el plazo leido no encaja, manda el importe (como el 10-oct).
        # Varios hitos iguales de la MISMA linea (50/50 sin plazo claro) = casa: se coge el primero libre al escribir.
        exactas = [o for o in opciones if C.cerca(o[2]['importe'], imp)] or por_importe
        if exactas:
            ids = {o[2]['id'] for o in exactas}
            lins = {o[1]['id'] for o in exactas}
            # (12-oct) lo que hace falta para escribirlo con C.escribir: la hoja, el hito y si hay hermanos iguales
            pri = sorted(exactas, key=lambda o: o[2]['orden'] or 0)[0]
            r.update({'veredicto': 'exacto' if len(ids) == 1 else 'exacto_varios', 'hito': pri[2]['id'], 'otros': sorted(ids),
                      'hoja': pri[0]['numero_hoja'], 'hoja_id': pri[0]['id']})
            causa('CASA AHORA' if len(ids) == 1 or len(lins) == 1 else 'VARIAS OPCIONES',
                  '; '.join('%s · %s %s' % (o[0]['numero_hoja'], o[2]['hito'], o[2]['importe']) for o in exactas[:4]))
            continue
        trozo = [o for o in opciones if o[2]['importe'] and imp < float(o[2]['importe']) - 1]
        if trozo:
            r.update({'veredicto': 'subhito', 'hito': trozo[0][2]['id'], 'hoja': trozo[0][0]['numero_hoja'], 'hoja_id': trozo[0][0]['id']})
            causa('SUBHITO', 'es un trozo de %s %s (%s)' % (trozo[0][2]['hito'], trozo[0][2]['importe'], trozo[0][0]['numero_hoja'])); continue
        causa('HITOS MAL GENERADOS', 'ningún hito es de ese plazo (%s) ni de ese importe' % x.get('plazo'))

    orden = ['causa', 'detalle', 'factura', 'fecha', 'estado_factura', 'cliente', 'importe', 'concepto', 'plazo', 'direccion', 'hoja_citada',
             'comunidad', 'hojas', 'lo_que_tiene_la_hoja', 'hitos', 'texto']
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'por que no casan'
    ws.append(orden)
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='2B2B2B')
    for r in sorted(filas, key=lambda r: (r['causa'], r['detalle'], r['factura'])):
        ws.append([r.get(k) for k in orden])
    for col, ancho in zip('ABCDEFGHIJKLMNOP', [22, 50, 20, 11, 11, 30, 10, 18, 11, 28, 11, 34, 40, 60, 50, 80]):
        ws.column_dimensions[col].width = ancho
    for fila in ws.iter_rows(min_row=2):
        for c in fila: c.alignment = Alignment(vertical='top', wrap_text=c.column_letter in 'BNOP')
    ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
    wb.save(SALIDA)
    json.dump(filas, open(SALIDA.replace('.xlsx', '.json'), 'w', encoding='utf-8'), ensure_ascii=False, default=str)
    for k_, v in cuenta.most_common(): print(f'{v:4}  {k_}')
    det = collections.Counter((r['causa'], re.sub(r'(HE|PD)-\d{4}-\d{4}|\d{4}-\d\d-\d\d|proyecto|css|df|subvencion|iee|informe', '…', r['detalle'])) for r in filas)
    print()
    for (c_, d), v in det.most_common(): print(f'{v:4}  {c_}: {d}')
    print('total', len(filas), '->', SALIDA)


if __name__ == '__main__':
    main()

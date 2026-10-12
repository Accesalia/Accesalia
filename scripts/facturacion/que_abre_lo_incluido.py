# -*- coding: utf-8 -*-
"""
QUE ABRE LO INCLUIDO SIN PRECIO (Monica, 12-oct-2026): "hay areas que se generan
desde las lineas cobrables y otras desde las NO cobrables". Las cobrables son
las lineas de facturacion (otro Excel: que_abre_cada_linea.xlsx, que ella esta
corrigiendo y NO se toca). Este saca lo NO cobrable: los conceptos de las hojas
firmadas que van incluidos, por bloque, para que diga que area abre cada uno.
SOLO LEE la base. Escribe un Excel NUEVO en Descargas.
"""
import collections, os, re, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from produccion import arrancar  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

SALIDA = os.path.join(os.path.expanduser("~"), "Downloads", "que_abre_lo_incluido.xlsx")

b = arrancar()
firmadas = {h["id"] for h in b.leer("hojas_encargo?select=id&estado=eq.devuelta_firmada")}
bloques = {x["id"]: x for x in b.leer("bloques?select=id,nombre,activo")}
con = [c for c in b.leer("conceptos_hoja?select=hoja_encargo_id,bloque_id,descripcion,incluido,importe") if c["hoja_encargo_id"] in firmadas and c["incluido"]]
cuenta = collections.Counter()
hojas = collections.defaultdict(set)
for c in con:
    k = bloques[c["bloque_id"]]["nombre"] if c["bloque_id"] in bloques else "(sin bloque) " + re.sub(r"\s+", " ", (c["descripcion"] or "").strip())[:70]
    cuenta[k] += 1
    hojas[k].add(c["hoja_encargo_id"])

wb = Workbook(); ws = wb.active; ws.title = "Lo incluido sin precio"
cab = ["Concepto incluido (bloque)", "Veces", "Hojas firmadas", "Área que abre", "Subárea", "Cuándo se abre", "Tus notas"]
ws.append(cab)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2B2B2B"); c.alignment = Alignment(wrap_text=True)
for k, n in cuenta.most_common():
    ws.append([k, n, len(hojas[k]), "", "", "", ""])
for col, a in zip("ABCDEFG", [60, 8, 14, 18, 22, 20, 40]):
    ws.column_dimensions[col].width = a
ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
wb.save(SALIDA)
print("Excel ->", SALIDA, "|", len(cuenta), "conceptos distintos en", len(firmadas), "hojas firmadas")

# -*- coding: utf-8 -*-
"""
QUE ABRE CADA LINEA FACTURABLE (Monica, 12-oct-2026): un Excel para que ella
decida, bloque a bloque, que AREA abre cada cosa que se cobra, con que apellido
(los codigos del 11-oct: DIR · OPP · AREA · apellido, docs/clave-actuacion-areas.md),
CUANDO se abre, y si va a FARO y a tecnico.

Va rellenado con una PROPUESTA (lo que dijo ella el 12-oct sobre FARO/tecnico y
lo del documento de areas) y una columna de dudas donde las dos cosas chocan.
Ella lo corrige; luego se carga como tabla editable.

SOLO LEE la base. Escribe el Excel en Descargas.
"""
import collections
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from produccion import arrancar  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

SALIDA = os.path.join(os.path.expanduser("~"), "Downloads", "que_abre_cada_linea.xlsx")

# bloque -> (area, apellido, cuando, duda)
PROPUESTA = {
    "REDACCIÓN DEL PROYECTO BÁSICO Y DE EJECUCIÓN": ("PRY", "ASC / SATE / CUB / ACC (según la obra)", "al cobrar", ""),
    "MODIFICACIÓN DE PROYECTO": ("PRY", "el del proyecto que modifica", "al cobrar", ""),
    "MEMORIA TÉCNICA SIMPLIFICADA": ("PRY", "memoria valorada", "al cobrar", ""),
    "CONSULTA URBANÍSTICA": ("PRY", "CON", "al cobrar", ""),
    "INFORME PERICIAL": ("PRY", "PER", "al cobrar",
                         "El documento del 11-oct lo pone en PRY (consume técnico), pero hoy has dicho que no va a técnico."),
    "MEDICIONES Y CIEGO": ("PRY", "parte del proyecto", "?",
                           "El 11-oct: parte de PRY. Hoy: no va ni a FARO ni a técnico. ¿Es parte del proyecto sin consumir técnico?"),
    "TOMA DE DATOS Y MODELADO 3D": ("—", "—", "no abre nada",
                                    "Hoy: es de la fase comercial, previa. El 11-oct se puso como parte de PRY."),
    "CERTIFICADO DE EFICIENCIA ENERGÉTICA": ("PRY", "CEE (parte del proyecto)", "?",
                                             "El 11-oct: 'el CEE es del proyecto'. Hoy: no va a técnico."),
    "TRAMITACIÓN DE LICENCIAS Y PERMISOS": ("PRY", "licencia (parte del proyecto)", "?", "¿Se abre al cobrar o al entregar el proyecto?"),
    "TRAMITACIÓN DE SUBVENCIONES": ("SUB", "por convocatoria presentada", "al cobrar", "Hoy: abre el área de subvenciones."),
    "COORDINACIÓN DE SEGURIDAD Y SALUD": ("OBR", "CSS", "al iniciar la obra", "Hoy: abre CSS cuando se inicia la obra."),
    "DIRECCIÓN FACULTATIVA": ("OBR", "DF", "¿al iniciar la obra?", ""),
    "FIN DE OBRA": ("OBR", "CFO", "?", ""),
    "INFORME DE EVALUACIÓN DEL EDIFICIO (IEE)": ("IEE", "IEE", "al cobrar", ""),
    "LIBRO DEL EDIFICIO": ("IEE", "LEE", "?", "El 11-oct: 'el LEE es del IEE'."),
    "GESTIÓN DE CAES": ("CAES", "", "?", "Lo lleva Ecobalance."),
    "SOLICITUD Y ANÁLISIS DE PRESUPUESTOS": ("3PP", "", "?", "Lo lleva Ecobalance."),
    "SOLICITUD DE FINANCIACIÓN": ("FIN", "", "?", "Apagado. Lo lleva Ecobalance."),
    "GESTIÓN DE SUBVENCIONES ACCESIBILIDAD": ("SUB", "", "", "Apagado: ya no se contrata."),
    "GESTIÓN DE SUBVENCIONES DE EFICIENCIA ENERGÉTICA": ("SUB", "", "", "Apagado: ya no se contrata."),
    "SATE CON CESION DE CAES": ("PRY + CAES", "SATE", "", "Apagado: ya no existe. Quedan 18 líneas antiguas."),
    "Honorarios de Proyecto:": ("PRY + CAES", "SATE + ASC", "",
                                "Apagado. Es 'SATE + ASCENSOR CON CESIÓN DE CAES' (el nombre se coló al importar). 1 línea: HE-2025-0077 / DAN-2025-037."),
}
FARO_TECNICO = {  # lo que dijo ella el 12-oct
    "REDACCIÓN DEL PROYECTO BÁSICO Y DE EJECUCIÓN", "MODIFICACIÓN DE PROYECTO", "MEMORIA TÉCNICA SIMPLIFICADA", "CONSULTA URBANÍSTICA",
}

# Las lineas sin bloque: por su texto, a que bloque irian (propuesta).
POR_TEXTO = [
    (r"COORDINACI|CSYS|\bCSS\b|RECURSO PREVENTIVO", "COORDINACIÓN DE SEGURIDAD Y SALUD"),
    (r"LIBRO DEL EDIFICIO", "LIBRO DEL EDIFICIO"),
    (r"\bCFO\b|FIN DE OBRA", "FIN DE OBRA"),
    (r"LICENCIA", "TRAMITACIÓN DE LICENCIAS Y PERMISOS"),
    (r"AYUDAS|SUBVENCION|NEXT GENERATION", "TRAMITACIÓN DE SUBVENCIONES"),
    (r"PROYECTO|ANTEPROYECTO", "REDACCIÓN DEL PROYECTO BÁSICO Y DE EJECUCIÓN"),
]

AREAS = [
    ("PRY", "Proyecto", "ASC, SATE, CUB, ACC… · CON (consulta urbanística) · PER (informe pericial) · memoria valorada", "Accesalia, equipo técnico",
     "Lo que consume tiempo de técnico. Visado, licencia, toma de datos, mediciones y CEE son partes de PRY (11-oct)."),
    ("OBR", "Obra", "DF, CSS, CFO", "Accesalia", "DF, CSS y CFO no son áreas: son partes de la obra."),
    ("SUB", "Subvención", "por convocatoria presentada", "Accesalia", ""),
    ("IEE", "IEE", "IEE, LEE", "Accesalia, fuera del equipo técnico", "El LEE es del IEE."),
    ("CAES", "CAES", "", "Ecobalance", ""),
    ("3PP", "3 presupuestos", "", "Ecobalance", ""),
    ("FIN", "Financiación", "", "Ecobalance", ""),
]

CAB = Font(bold=True, color="FFFFFF")
FONDO_CAB = PatternFill("solid", fgColor="2B2B2B")
AMARILLO = PatternFill("solid", fgColor="FFF5CC")  # lo que tiene que mirar ella
GRIS = PatternFill("solid", fgColor="EEEEEE")       # apagado


def hoja(wb, titulo, cabecera, filas, anchos, pintar=None):
    ws = wb.create_sheet(titulo)
    ws.append(cabecera)
    for c in ws[1]:
        c.font, c.fill = CAB, FONDO_CAB
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for f in filas:
        ws.append(f)
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = a
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if pintar:
            relleno = pintar([c.value for c in fila])
            if relleno:
                for c in fila:
                    c.fill = relleno
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    b = arrancar()
    bloques = b.leer("bloques?select=id,codigo,nombre,activo,va_a_faro,va_a_tecnico,orden&order=orden.asc")
    lineas = b.leer("lineas_facturacion?select=id,bloque_id,descripcion,importe,hoja:hoja_encargo_id(numero_hoja,estado)")
    por_bloque = collections.Counter(l["bloque_id"] for l in lineas)

    wb = Workbook()
    wb.remove(wb.active)

    # 1 · un bloque por fila: lo que hay que decidir
    filas = []
    for x in sorted(bloques, key=lambda x: (not x["activo"], -por_bloque.get(x["id"], 0))):
        area, apellido, cuando, duda = PROPUESTA.get(x["nombre"], ("?", "?", "?", "Sin propuesta"))
        sí = x["nombre"] in FARO_TECNICO
        filas.append([
            x["nombre"], "sí" if x["activo"] else "no (apagado)", por_bloque.get(x["id"], 0),
            area, apellido, cuando, "sí" if sí else "no", "sí" if sí else "no", duda, "",
        ])
    hoja(
        wb, "Qué abre cada bloque",
        ["Bloque (lo que se cobra)", "Contratable", "Líneas", "Área", "Subárea / apellido", "Cuándo se abre",
         "¿Va a FARO?", "¿Va a técnico?", "Duda o choque", "Tus notas"],
        filas, [44, 13, 8, 10, 30, 20, 11, 12, 55, 35],
        pintar=lambda v: GRIS if str(v[1]).startswith("no") else (AMARILLO if v[8] else None),
    )

    # 2 · las lineas sin bloque, agrupadas por su texto, con el bloque que les tocaria
    sin = [l for l in lineas if not l["bloque_id"]]
    grupos = collections.defaultdict(list)
    for l in sin:
        t = re.sub(r"\s+", " ", (l["descripcion"] or "(sin texto)").strip())
        grupos[t].append(l)
    filas = []
    for t, ls in sorted(grupos.items(), key=lambda kv: -len(kv[1])):
        propuesto = next((blq for pat, blq in POR_TEXTO if re.search(pat, t.upper())), "?")
        hojas = sorted({(l.get("hoja") or {}).get("numero_hoja") or "—" for l in ls})
        filas.append([t, len(ls), round(sum(float(l["importe"] or 0) for l in ls), 2), propuesto, ", ".join(hojas[:8]) + (" …" if len(hojas) > 8 else ""), ""])
    hoja(
        wb, f"Líneas sin bloque ({len(sin)})",
        ["Texto de la línea", "Líneas", "Importe total", "Bloque que le tocaría (propuesta)", "Hojas (algunas)", "Tus notas"],
        filas, [60, 8, 14, 44, 45, 35],
        pintar=lambda v: AMARILLO if v[3] == "?" else None,
    )

    # 3 · la leyenda de las areas (11-oct)
    hoja(
        wb, "Las áreas (11-oct)",
        ["Código", "Área", "Apellidos", "Quién", "Notas"],
        [list(a) for a in AREAS], [9, 16, 60, 30, 60],
    )
    ws = wb["Las áreas (11-oct)"]
    ws.append([])
    ws.append(["Código de cada trabajo:", "DIR-00412 · DAN-2026-060 · ÁREA · apellido   (hasta el área, común; el apellido lo decide cada área)"])

    wb.save(SALIDA)
    print("Excel ->", SALIDA, "|", len(bloques), "bloques,", len(sin), "líneas sin bloque en", len(grupos), "textos")


if __name__ == "__main__":
    main()

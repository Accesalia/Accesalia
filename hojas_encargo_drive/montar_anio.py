# Por AÑOS (Monica: 'solo hacemos de año en año'): python montar_anio.py 2024
# Claras por concepto y fecha; dudosas y sin calle, leyendo (instrucciones_dudosas.md).
# Antes: tanda 2. Una fila por HOJA (enviada),
# con su firmada. Mismas reglas que la tanda 1 (normalizar) y lo resuelto leyendo.
import csv, json, glob, os, re, collections, sys
A = sys.argv[1]
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

exec(open('cruce.py', encoding='utf-8').read().split('fuentes=collections')[0])  # clave(), fecha(), norm()
src = open('montar_excel.py', encoding='utf-8').read()
exec(src[src.index('# REGLAS DE MONICA'):src.index('def cabecera(n):')])         # normalizar()

lote3 = json.load(open(f'anio{A}_parejas.json', encoding='utf-8'))
idx = json.load(open(f'anio{A}_indice.json', encoding='utf-8'))
ext = {}
for f in glob.glob(f'extraido{A}*.jsonl'):
    for l in open(f, encoding='utf-8'):
        if l.strip():
            d = json.loads(l); ext[d['n']] = d
dud = {}
for f in sorted(glob.glob(f'dud{A}_resueltas_*.jsonl')):
    for l in open(f, encoding='utf-8'):
        d = json.loads(l); dud[d['d']] = d
res = {}
if os.path.exists('resoluciones_' + A + '.json'):
    res = {k: v for k, v in json.load(open('resoluciones_' + A + '.json', encoding='utf-8')).items() if not k.startswith('_')}

def cabecera(n):
    p = f'textos{A}/{n}.txt'
    if not os.path.exists(p): return {}
    l = open(p, encoding='utf-8').readline()
    if not l.startswith('#DRIVE_ID='): return {}
    return dict(x.lstrip('#').split('=', 1) for x in l.strip().split('\t') if '=' in x)

PROBLEMA = r'duplicad|dos hojas|otra propuesta|no cuadra|copiad|errata|aunque el edificio|no es una hoja|ilegible|tachad|alternativ|anulad|otra comunidad|de otra'

# Las firmadas, con su(s) enviada(s) y como se decidio
filas = []
for x in lote3:
    p = {'firmada': x['firmada'], 'recibida': x['recibida'] or ''}
    if x['caso'] in ('dudosa', 'sin_clave'):
        dd = dud[x['d']]; filas.append((p, dd['decision'], dd['elegidas'], dd))
    elif x['caso'] == 'no_es_hoja':
        filas.append((p, 'no es hoja', [], {'motivo': 'por el nombre: documento de facturación o similar'}))
    else:
        filas.append((p, 'pareja', x['elegidas'], None))

usos = collections.Counter(e for _, dec, el, _ in filas if dec == 'pareja' for e in el)


# Quien GENERO la hoja segun el Excel de las automaticas (columna W de seleccion_encargos).
# Monica, 6-oct: "poco fiable"; CG Carlos Garcia, CS Carlos Sepulveda, Alejandra por
# Daniel, Alvaro ni lo pone. SOLO sirve de cotejo con el comercial asignado de la opp.
def _quien_genero():
    import openpyxl, re
    out = {}
    try:
        ws = openpyxl.load_workbook('seleccion_encargos.xlsx', read_only=True, data_only=True)['seleccion_encargos']
        for r in ws.iter_rows(min_row=2, values_only=True):
            m = re.search(r'id=([\w-]+)', str(r[21] or ''))
            if m and r[22]: out[m.group(1)] = str(r[22]).strip()
    except Exception:
        pass
    return out
QUIEN = _quien_genero()

wb = Workbook(); ws = wb.active; ws.title = 'Cruce'
cols = ['N', 'Estado propuesto', 'Por qué', 'Fecha recibida (firmada)', 'Fecha de la hoja (enviada)',
        'Dirección (en la hoja)', 'A quién', 'Qué se hace', 'Sin casar (qué se hace)', 'Conceptos',
        'Total base (sin IVA)', 'Forma de pago', 'Firmada (fichero)', 'Enviada (documento)', 'Enviada: carpeta',
        'Enviada: última modificación', 'Cómo se emparejó', 'Notas de la lectura', 'Generada por (Excel, solo cotejo)', 'Código HE', 'Opp', 'TU DECISIÓN']
ws.append(cols)
wc = wb.create_sheet('Conceptos')
wc.append(['N', 'Dirección', 'Concepto (texto de la hoja)', 'Bloque del catálogo', 'Importe base', '% a éxito', 'Incluido (no se cobra aparte)', 'Forma de pago'])
cuenta = collections.Counter()

def fila(N, estado, motivos, frec, d, firm, env, como, nota_extra=None):
    cab = cabecera(idx.get(env, '')) if env else {}
    conc = d.get('conceptos') or []
    txt = '\n'.join(f"{c['bloque']}: " + (f"{c['importe']:g} €" if c.get('importe') is not None else f"{c['pct_exito']:g}% a éxito" if c.get('pct_exito') else 'incluido' if c.get('incluido') else 'sin precio propio') for c in conc)
    notas = '; '.join(x for x in (d.get('rarezas'), nota_extra) if x)
    ws.append([N, estado, '; '.join(motivos), frec, d.get('fecha_hoja'), d.get('direccion'), d.get('a_quien'),
               ', '.join(d.get('que_se_hace') or []), ', '.join(d.get('que_se_hace_sin_casar') or []), txt,
               d.get('total_base'), d.get('forma_pago_general'), firm.split('/')[-1],
               (cab.get('TITULO') or (env or '').split('/')[-1]), (env or '').split('/')[0], cab.get('MODIFICADO', '')[:10],
               como, notas or None, QUIEN.get(cab.get('DRIVE_ID','')), 'se asigna al final, con todas', '', ''])
    for c in conc:
        wc.append([N, d.get('direccion'), c.get('texto'), c.get('bloque'), c.get('importe'), c.get('pct_exito'), 'sí' if c.get('incluido') else '', c.get('forma_pago')])
    cuenta[estado] += 1

for k, (p, dec, el, dd) in enumerate(filas):
    frec = p['recibida']; firm = p['firmada']; base = lote3[k]['d']
    como = {'pareja': ('leyendo la firmada: ' + dd['motivo']) if dd else 'por concepto y fecha del nombre',
            'datos de la firmada': 'sin enviada que case; datos leídos de la firmada: ' + (dd or {}).get('motivo', ''),
            'no es hoja': 'no es una hoja de encargo: ' + (dd or {}).get('motivo', ''), 'no se puede': (dd or {}).get('motivo', '')}[dec]
    if dec == 'no es hoja':
        fila(base, 'APARTAR', ['no es una hoja: documento de facturación o similar'], frec, {}, firm, None, como); continue
    if dec in ('no se puede', 'no es hoja') and res.get(base, {}).get('estado') in ('APARTAR', 'ANULADA'):
        fila(base, res[base]['estado'], ['RESUELTO LEYENDO: ' + res[base].get('nota', '')], frec, {}, firm, None, como); continue
    if dec == 'no se puede':
        fila(base, 'PREGUNTA', ['PARA TI: ' + dd['motivo']], frec, {}, firm, None, como); continue
    if dec == 'datos de la firmada':
        d = {'fecha_hoja': dd.get('fecha_emision_firmada'), 'a_quien': dd.get('a_quien'), 'conceptos': [],
             'rarezas': None}
        fila(base, 'OK', [], frec, d, firm, None, como, 'Importes en la firmada: ' + (dd.get('importes_firmada') or '?')); continue
    for j, env in enumerate(el):
        N = base + (chr(97 + j) if len(el) > 1 else '')
        if res.get(N, {}).get('enviada'): env = res[N]['enviada']
        n = idx.get(env); d = normalizar(ext[n]) if n in ext else {}
        motivos = []
        if not d: motivos.append('enviada sin leer')
        fh = d.get('fecha_hoja')
        if dd and dd.get('fecha_emision_firmada') and fh and dd['fecha_emision_firmada'] != fh:
            motivos.append(f"la firmada es del {dd['fecha_emision_firmada']} y la enviada del {fh}")
        if fh and frec and fh > frec: motivos.append('la hoja es POSTERIOR a la firmada')
        cab = cabecera(n) if n else {}
        mod = cab.get('MODIFICADO', '')[:10]
        if mod and frec and mod > frec and not dd: motivos.append('el documento se tocó DESPUÉS de llegar firmada')
        if usos[env] > 1: motivos.append(f'la misma enviada sirve a {usos[env]} firmadas')
        if d.get('rarezas') and re.search(PROBLEMA, d['rarezas'], re.I): motivos.append('la lectura avisa de algo raro')
        if any(c.get('bloque') == 'SIN CASAR' for c in d.get('conceptos') or []) or d.get('que_se_hace_sin_casar'):
            motivos.append('concepto sin casar con el catálogo')
        estado = 'REVISAR' if motivos else 'OK'
        r = res.get(N)
        if r:
            estado = r['estado']
            motivos = [('PARA TI: ' if estado == 'PREGUNTA' else 'RESUELTO LEYENDO: ') + r['nota']] if r.get('nota') else []
            for campo in ('fecha_hoja', 'total_base', 'a_quien', 'que_se_hace', 'conceptos', 'que_se_hace_sin_casar', 'forma_pago_general'):
                if campo in r: d = dict(d); d[campo] = r[campo]
        fila(N, estado, motivos, frec, d, firm, env, como)

color = {'OK': 'DCEFC8', 'REVISAR': 'FFF2CC', 'PREGUNTA': 'FFD966', 'APARTAR': 'F4CCCC', 'ANULADA': 'D9D9D9'}
for hoja, anchos in ((ws, [7, 11, 40, 12, 12, 30, 18, 20, 18, 42, 12, 40, 45, 45, 22, 14, 45, 50, 12, 14, 10, 30]),
                     (wc, [7, 30, 45, 30, 12, 9, 12, 45])):
    for j, a in enumerate(anchos, 1): hoja.column_dimensions[get_column_letter(j)].width = a
    for c in hoja[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='2B2B2B')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    hoja.freeze_panes = 'B2'; hoja.auto_filter.ref = hoja.dimensions
    for f in hoja.iter_rows(min_row=2):
        for c in f: c.alignment = Alignment(wrap_text=True, vertical='top')
for f in ws.iter_rows(min_row=2):
    f[1].fill = PatternFill('solid', fgColor=color[f[1].value]); f[21].fill = PatternFill('solid', fgColor='EAF4FF')
wb.save(f'../docs/cruce_hojas_{A}.xlsx')
print(dict(cuenta), 'filas', ws.max_row - 1, 'extraidas', len(ext))

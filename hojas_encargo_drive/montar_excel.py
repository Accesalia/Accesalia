# Monta la tabla de cruce enviada <-> firmada para que Monica la revise.
import csv, json, glob, os, re, collections, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

exec(open('cruce.py', encoding='utf-8').read().split('fuentes=collections')[0])  # clave(), fecha(), norm()

tanda = list(csv.reader(open('tanda1.tsv', encoding='utf-8'), delimiter='\t'))
ext = {}
for f in glob.glob('extraido_*.jsonl'):
    for l in open(f, encoding='utf-8'):
        if l.strip():
            d = json.loads(l); ext[d['n']] = d
res = {k: v for k, v in json.load(open('resoluciones_tanda1.json', encoding='utf-8')).items() if not k.startswith('_')}
monday = json.load(open('monday_candidatas.json', encoding='utf-8'))

# Otras versiones del mismo encargo en las carpetas de fuentes ("Modificado ...",
# "modificado tras reunion ...") : misma calle contenida y mismo numero.
fuentes = []
for row in csv.reader(open('fuentes.tsv', encoding='utf-8'), delimiter='\t'):
    nom = row[2].split('/')[-1]; k = clave(nom)
    if k: fuentes.append((k, nom))

def otras_versiones(k, elegido):
    out = []
    for k2, nom in fuentes:
        if k2[1] == k[1] and k[0] <= k2[0] and nom != elegido and re.search(r'modific|v ?2|version', nom, re.I):
            out.append(nom)
    return out


# REGLAS DE MONICA (6-oct-2026), aplicadas al montar; lo extraido en bruto no se toca:
#  1. Subvenciones de accesibilidad y de eficiencia energetica son "subvenciones sin mas".
#  2. La "Documentacion Tecnica Anexa" (IEE, CEE, LEE, comparativa...) es lo que ahora se
#     llama doc tecnica: el importe FIJO de la linea de subvencion (el % es lo de a exito).
SUBV = 'TRAMITACION SUBVENCIONES'
ES_DOC = re.compile(r'doc(umentaci[oó]n)?\.?\s*t[eé]cnica', re.I)
def normalizar(d):
    conc = [dict(c) for c in d.get('conceptos') or []]
    for c in conc:
        if c.get('bloque', '').startswith(SUBV): c['bloque'] = SUBV
    docs = [c for c in conc if c.get('bloque') == 'SIN CASAR' and ES_DOC.search(c.get('texto') or '')]
    subv = next((c for c in conc if c.get('bloque') == SUBV), None)
    for c in docs:
        if subv is not None:
            if c.get('importe') is not None and subv.get('importe') is None:
                subv['importe'] = c['importe']
            conc.remove(c)          # su contenido va dentro de la subvencion
        else:
            c['bloque'] = SUBV; subv = c
    # 3. (Monica, 6-oct) "mediciones y presupuesto ciego" es el bloque MEDICIONES Y CIEGO.
    for c in conc:
        if c.get('bloque') == 'SIN CASAR' and re.search(r'medici', c.get('texto') or '', re.I) and re.search(r'ciego', c.get('texto') or '', re.I):
            c['bloque'] = 'MEDICIONES Y CIEGO'
    # 5. CAES (Monica, 6-oct): "¿la cobramos? linea que se cobra. ¿No la cobramos? linea
    #    incluida. ¿No se menciona? no sale." Un descuento por cesion = no se cobra.
    ES_CAES = re.compile(r'(?<![A-Za-z])CAES?(?![A-Za-z])|AHORRO ENERG', re.I)
    for c in conc:
        if c.get('bloque') == 'SIN CASAR' and ES_CAES.search(c.get('texto') or ''):
            if (c.get('importe') or 0) > 0:
                c['bloque'] = 'GESTION DE CAES'; c['incluido'] = False
            else:
                c['bloque'] = 'GESTION DE CAES'; c['importe'] = None; c['incluido'] = True
    menciona = any(ES_CAES.search((c.get('texto') or '') + ' ' + (c.get('bloque') or '')) for c in conc)
    if menciona and not any(c.get('bloque') == 'GESTION DE CAES' for c in conc):
        conc.append({'texto': 'CAES (no se cobran aparte)', 'bloque': 'GESTION DE CAES', 'importe': None,
                     'pct_exito': None, 'forma_pago': None, 'incluido': True})
    # 7. Next Generation (Monica, 6-oct): el "adelanto a cuenta del importe total" NO es un
    #    concepto: es la forma de pago (el resto, solo si se concede la subvencion).
    for c in [c for c in conc if re.search(r'adelanto', c.get('texto') or '', re.I) and c.get('bloque') == 'SIN CASAR']:
        conc.remove(c)
        # Si la forma de pago ya cuenta el adelanto, se queda como esta; si no, se le pone.
        if c.get('importe') and not re.search(r'adelanto', d.get('forma_pago_general') or '', re.I):
            imp = f"{c['importe']:,.0f}".replace(',', '.')
            d = dict(d); d['forma_pago_general'] = f"Adelanto a cuenta {imp} € + IVA a la firma; el resto, solo si se concede la subvención." + (' ' + d['forma_pago_general'] if d.get('forma_pago_general') else '')
    # 6. Conceptos sueltos que son lo que son (leidos en las tandas 1 y 2).
    for c in conc:
        if c.get('bloque') != 'SIN CASAR': continue
        t = c.get('texto') or ''
        if re.search(r'inicio de obra', t, re.I): c['bloque'] = 'DF'   # hito de la DF con precio propio
        elif re.search(r'libro del edificio', t, re.I): c['bloque'] = 'LEE'   # 'Libro del Edificio + IEE' con precio: el LEE (IEE dentro)
        elif re.search(r'(?<![A-Za-z])LEE(?![A-Za-z])', t): c['bloque'] = 'LEE'   # 'LEE + IEE' con un precio: el LEE
        elif re.search(r'(?<![A-Za-z])IEE(?![A-Za-z])', t): c['bloque'] = 'IEE'   # 'CEE e IEE' con un precio
        elif re.search(r'(?<![A-Za-z])DF(?![A-Za-z])|DIRECCI', t, re.I): c['bloque'] = 'DF'   # 'DF + CSS' con un precio
        elif re.search(r'proyecto', t, re.I): c['bloque'] = 'REDACCION PROYECTO'
        elif re.search(r'inspecci', t, re.I): c['bloque'] = 'INFORME PERICIAL'
        elif re.search(r'asesoramiento.*(ayuda|subvenc)', t, re.I): c['bloque'] = 'TRAMITACION SUBVENCIONES'; c['incluido'] = True
        elif re.search(r'consulta urban', t, re.I): c['bloque'] = 'CONSULTA URBANISTICA'
        elif re.search(r'licencia', t, re.I): c['bloque'] = 'TRAMITACION LICENCIAS'
        elif re.search(r'visado|canal de isabel|alta de(l)? agua', t, re.I): c['incluido'] = True; c['bloque'] = 'TRAMITACION LICENCIAS'   # tramites sin bloque propio
    # 4. Que se hace: elevador = plataforma elevadora = Plataforma; portal = Accesibilidad portal.
    que = list(d.get('que_se_hace') or []); sin = []
    for t in d.get('que_se_hace_sin_casar') or []:
        if re.search(r'elevador', t, re.I): que.append('Plataforma')
        elif re.search(r'portal', t, re.I): que.append('Accesibilidad portal')
        elif re.search(r'medici.*ciego', t, re.I): pass
        else: sin.append(t)
    d = dict(d); d['conceptos'] = conc; d['que_se_hace'] = list(dict.fromkeys(que)); d['que_se_hace_sin_casar'] = sin
    return d

def cabecera(n):
    p = f'textos/{n}.txt'
    if not os.path.exists(p): return {}
    l = open(p, encoding='utf-8').readline()
    if not l.startswith('#DRIVE_ID='): return {}
    return dict(x.lstrip('#').split('=', 1) for x in l.strip().split('\t') if '=' in x)

def codigo_previo(n):
    p = f'textos/{n}.txt'
    if not os.path.exists(p): return None
    m = re.findall(r'(?<![0-9])(20[0-9][0-9]/[0-9][0-9]/[0-9]{3})(?![0-9])', open(p, encoding='utf-8').read())
    return ', '.join(dict.fromkeys(m)) or None

def eur(x):
    return None if x is None else float(x)


# Quien GENERO la hoja segun el Excel de las automaticas (columna W de seleccion_encargos).
# Monica, 6-oct: "poco fiable"; CG Carlos Garcia, CS Carlos Sepulveda, Alejandra por
# Daniel, Alvaro ni lo pone. SOLO sirve de cotejo con el comercial asignado de la opp.
def _quien_genero():
    import openpyxl, re
    out = {}
    try:
        ws = openpyxl.load_workbook('seleccion_encargos.xlsx', read_only=True, data_only=True)['seleccion_encargos']
        for r in ws.iter_rows(min_row=2, values_only=True):
            m = re.search(r'id=([\w-]+)', str(r[21] or ''))
            if m and r[22]: out[m.group(1)] = str(r[22]).strip()
    except Exception:
        pass
    return out
QUIEN = _quien_genero()

wb = Workbook()
ws = wb.active; ws.title = 'Cruce'
cols = ['N', 'Estado propuesto', 'Por qué', 'Fecha recibida (firmada)', 'Fecha de la hoja (enviada)',
        'Dirección (en la hoja)', 'A quién', 'Qué se hace', 'Sin casar (qué se hace)', 'Conceptos',
        'Total base (sin IVA)', 'Forma de pago',
        'Firmada (fichero)', 'Enviada (documento)', 'Enviada: carpeta', 'Enviada: última modificación',
        'Otras versiones en la carpeta', 'Notas de la lectura', 'Generada por (Excel, solo cotejo)', 'Código HE', 'Opp', 'TU DECISIÓN']
ws.append(cols)
wc = wb.create_sheet('Conceptos')
wc.append(['N', 'Dirección', 'Concepto (texto de la hoja)', 'Bloque del catálogo', 'Importe base', '% a éxito', 'Incluido (no se cobra aparte)', 'Forma de pago'])

cuenta = collections.Counter()
usos = collections.Counter(e for _, _, e in tanda)
for i, (frec, firm, env) in enumerate(tanda):
    n = f'{i:03d}'; d = normalizar(ext[n]) if n in ext else {}; cab = cabecera(n)
    firm_nom = firm.split('/')[-1]; env_nom = cab.get('TITULO') or env.split('/')[-1]
    k = clave(firm_nom)
    motivos = []
    if re.search(r'no vale|anulad|cancelad', firm_nom, re.I): estado = 'APARTAR'; motivos.append('la firmada se llama "no vale/anulado"')
    elif not d or d.get('rarezas') == 'sin texto': estado = 'APARTAR'; motivos.append('enviada sin texto')
    else: estado = None
    fh = d.get('fecha_hoja')
    if fh and fh > frec: motivos.append('la hoja es POSTERIOR a la firmada')
    # La vigencia de 3 meses NO se tiene en cuenta (Monica): no se avisa del tiempo entre hoja y firma.
    mod = cab.get('MODIFICADO', '')[:10]
    if mod and mod > frec: motivos.append('el documento se tocó DESPUÉS de llegar firmada')
    ov = otras_versiones(k, env.split('/')[-1]) if k else []
    if ov: motivos.append('hay otra versión en la carpeta')
    if usos[env] > 1: motivos.append(f'la misma enviada sirve a {usos[env]} firmadas')
    cand = monday.get(n, [])
    PROBLEMA = r'duplicad|dos hojas|otra propuesta|no cuadra|copiad|errata|aunque el edificio|no es una hoja|ilegible|tachad|alternativ|anulad|otra comunidad|de otra'
    if d.get('rarezas') and re.search(PROBLEMA, d['rarezas'], re.I): motivos.append('la lectura avisa de algo raro')
    if any(c.get('bloque') == 'SIN CASAR' for c in d.get('conceptos') or []) or d.get('que_se_hace_sin_casar'):
        motivos.append('concepto sin casar con el catálogo')
    if not estado: estado = 'REVISAR' if motivos else 'OK'
    # Lo resuelto leyendo manda sobre lo automatico.
    r = res.get(n)
    if r:
        estado = r['estado']
        motivos = ['RESUELTO LEYENDO: ' + r['nota']] if r.get('nota') else []
        if r['estado'] == 'PREGUNTA': motivos = ['PARA TI: ' + r['nota']]
        for campo in ('fecha_hoja', 'total_base', 'a_quien', 'que_se_hace', 'conceptos', 'que_se_hace_sin_casar', 'forma_pago_general'):
            if campo in r: d = dict(d); d[campo] = r[campo]
        fh = d.get('fecha_hoja')
    cuenta[estado] += 1
    conc = d.get('conceptos') or []
    txt_conc = '\n'.join(
        f"{c['bloque']}: " + (f"{c['importe']:g} €" if c.get('importe') is not None else f"{c['pct_exito']:g}% a éxito" if c.get('pct_exito') else 'incluido' if c.get('incluido') else 'sin precio propio')
        for c in conc)
    mon = '\n'.join(f"{c['desc']} · {c['estado']} · {c['fecha']} · id {c['id'][:8]}" for c in cand)
    ws.append([n, estado, '; '.join(motivos), frec, fh, d.get('direccion'), d.get('a_quien'),
               ', '.join(d.get('que_se_hace') or []), ', '.join(d.get('que_se_hace_sin_casar') or []),
               txt_conc, eur(d.get('total_base')), d.get('forma_pago_general'),
               firm_nom, env_nom, env.split('/')[0], mod, '\n'.join(ov), d.get('rarezas'), QUIEN.get(cab.get('DRIVE_ID','')),
               'se asigna al final, con todas', '', ''])
    for c in conc:
        wc.append([n, d.get('direccion'), c.get('texto'), c.get('bloque'), eur(c.get('importe')), c.get('pct_exito'), 'sí' if c.get('incluido') else '', c.get('forma_pago')])

# formato
color = {'OK': 'DCEFC8', 'REVISAR': 'FFF2CC', 'PREGUNTA': 'FFD966', 'APARTAR': 'F4CCCC', 'ANULADA': 'D9D9D9'}
for hoja, anchos in ((ws, [5, 11, 40, 12, 12, 30, 18, 20, 18, 42, 12, 40, 45, 45, 22, 14, 35, 50, 12, 14, 10, 30]),
                     (wc, [5, 30, 45, 30, 12, 9, 12, 45])):
    for j, a in enumerate(anchos, 1):
        hoja.column_dimensions[get_column_letter(j)].width = a
    for c in hoja[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='2B2B2B')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    hoja.freeze_panes = 'B2'; hoja.auto_filter.ref = hoja.dimensions
    for fila in hoja.iter_rows(min_row=2):
        for c in fila: c.alignment = Alignment(wrap_text=True, vertical='top')
for fila in ws.iter_rows(min_row=2):
    fila[1].fill = PatternFill('solid', fgColor=color[fila[1].value])
    fila[21].fill = PatternFill('solid', fgColor='EAF4FF')
salida = '../docs/cruce_hojas_tanda1.xlsx'
wb.save(salida)
print(dict(cuenta), 'extraidas', len(ext), '->', salida)

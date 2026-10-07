# Prepara lo que hay que LEER de un año para su manifiesto, en tres pasos (en este orden, porque
# cada uno usa lo resuelto en el anterior):
#   python candidatos_anio.py 2025 pv          -> posibles_versiones_2025.json + lote_pv_2025_k.json
#   python candidatos_anio.py 2025 su          -> candidatas_sustitucion_2025.json + lote_su_2025_k.json
#   python candidatos_anio.py 2025 sinfirmar   -> sinfirmar_2025.json + lote_u_2025_k.json (+ textos de Word/PDF del disco)
# Los agentes leen con instrucciones_versiones.md / instrucciones_sustitucion.md / instrucciones_sinfirmar.md
# y escriben pv_resueltas_<A>_k.jsonl / su_resueltas_<A>_k.jsonl / extraido<A>u_k.jsonl.
import sys, json, collections, os, subprocess
A, PASO = sys.argv[1], sys.argv[2]
LOTES = int(sys.argv[3]) if len(sys.argv) > 3 else 3
sys.argv = ['x', A]
exec(open('manifiesto_anio.py', encoding='utf-8').read().split('# 7 · el Excel')[0])

def reparte(lista, prefijo):
    for k in range(LOTES):
        json.dump(lista[k::LOTES], open(f'{prefijo}_{k}.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(PASO, len(lista), [len(lista[k::LOTES]) for k in range(LOTES)])

idd = lambda h: id_por_titulo.get(key_nom(h['titulo'])) or (h['xl'] or {}).get('id') if (h['xl'] or {}).get('id') or id_por_titulo.get(key_nom(h['titulo'])) else None
ficha = lambda h: {'titulo': h['titulo'], 'fecha': h['fecha'], 'id': idd(h), 'ruta': (h['enviados'] or [''])[0], 'firmada': bool(h['firmadas'])}

if PASO == 'pv':
    G = collections.defaultdict(list)
    for h in hojas:
        if h.get('version_de') or not h['k'] or not h['et'] or not h['enviados']: continue
        G[(h['k'][0], h['k'][1], frozenset(h['et']))].append(h)
    grupos = []
    for hs in G.values():
        if len({h['fecha'] for h in hs}) < 2: continue
        hs.sort(key=lambda h: h['fecha'])
        grupos.append({'g': f'PV{len(grupos):02d}', 'hojas': [ficha(h) for h in hs]})
    json.dump(grupos, open(f'posibles_versiones_{A}.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    reparte(grupos, f'lote_pv_{A}')

elif PASO == 'su':
    por = collections.defaultdict(list)
    for g in filas:
        if g[-1]['k']: por[(g[-1]['k'][0], g[-1]['k'][1])].append(g)
    grupos = []
    for gs in por.values():
        gs.sort(key=lambda g: g[0]['fecha'])
        for i, g in enumerate(gs):
            h = g[-1]
            if any(x['firmadas'] for x in g) or h.get('sustituida_por') or not h['enviados']: continue
            despues = [x for x in gs[i + 1:] if x[0]['fecha'] > h['fecha'] and not x[-1].get('sustituida_por')]
            cubre = set().union(*[x[-1]['et'] for x in despues]) if despues else set()
            if despues and h['et'] and h['et'] <= cubre:
                grupos.append({'g': f'SU{len(grupos):02d}', 'antigua': ficha(h), 'posteriores': [ficha(x[-1]) for x in despues]})
    json.dump(grupos, open(f'candidatas_sustitucion_{A}.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    reparte(grupos, f'lote_su_{A}')

elif PASO == 'sinfirmar':
    # Los Word/PDF del disco se leen aqui mismo (texto a textos<A>u/); los Google Docs, los agentes.
    import docx
    B = 'G:/Mi unidad/MONICA ACCESALIA/PRESUPUESTOS/'
    os.makedirs(f'textos{A}u', exist_ok=True)
    lista = []
    for g in filas:
        if any(h['firmadas'] for h in g): continue
        for h in g:
            if not h['enviados']: continue
            lista.append(ficha(h))
    for i, x in enumerate(lista):
        x['n'] = f'U{A[2:]}_{i:03d}'
        ruta = x['ruta']; ext = ruta.rsplit('.', 1)[-1].lower(); out = f'textos{A}u/{x["n"]}.txt'
        x['leido_del_disco'] = False
        try:
            if ext == 'docx':
                d = docx.Document(B + ruta); p = [q.text for q in d.paragraphs]
                for t in d.tables:
                    for r in t.rows: p.append(' | '.join(c.text.strip() for c in r.cells))
                open(out, 'w', encoding='utf-8').write('\n'.join(p)); x['leido_del_disco'] = True
            elif ext == 'pdf':
                open(out, 'w', encoding='utf-8').write(subprocess.run(['pdftotext', '-layout', B + ruta, '-'], capture_output=True).stdout.decode('utf-8', 'ignore')); x['leido_del_disco'] = True
        except Exception as e:
            print('no se pudo leer', ruta, e)
    json.dump(lista, open(f'sinfirmar_{A}.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    disco = [x for x in lista if x['leido_del_disco']]; drive = [x for x in lista if not x['leido_del_disco']]
    open(f'lote_u_{A}_disco.txt', 'w').write(' '.join(x['n'] for x in disco))
    print('del disco (solo extraer):', len(disco), '| google docs (leer y extraer):', len(drive))
    reparte(drive, f'lote_u_{A}')

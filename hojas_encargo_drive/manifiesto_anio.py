# MANIFIESTO DE VOLCADO de un año (Monica, 7-oct-2026): una fila por HOJA, con su codigo
# provisional por fecha, sus versiones y TODOS sus ficheros (enviados y firmados). Incluye las
# enviadas SIN firmar: "una señal muy fuerte: haz algo con esto, persiguelo o cierralo".
#   python manifiesto_anio.py 2026
#
# Universo: los documentos enviados del año en las carpetas (fecha en el nombre), mas las
# firmadas cuyos datos salieron de la propia firmada (sin enviada en la carpeta).
# Versiones: SOLO con señal explicita en el titulo (Modificado/MODF/REV/modificado tras reunion)
# sobre la misma direccion y concepto. El parecido sin señal se marca "posible version", no se junta.
import csv, re, sys, json, datetime, collections, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
exec(open('cruce.py', encoding='utf-8').read().split('fuentes=collections')[0])   # clave(), fecha(), norm()
exec(open('elegir_tanda2.py', encoding='utf-8').read().split('fuentes = []')[0].split('ETIQUETAS = ')[0])
src = open('elegir_tanda2.py', encoding='utf-8').read()
exec(src[src.index('ETIQUETAS = '):src.index('fuentes = []')])                    # etiquetas(), NO_ES_HOJA
A = sys.argv[1]
B = 'G:/Mi unidad/MONICA ACCESALIA/PRESUPUESTOS/'
# "MODF ASC" / "MODIF.ASC" = modificacion de ascensor, NO un modificado. Version solo si el titulo
# EMPIEZA por Modificado/Modificada/MODF, o dice "modificado tras" / "modificacion de condiciones", o "REV n".
_MOD = re.compile(r'^\s*(MODIFICAD[OA]|MODF)\b|MODIFICADO TRAS|MODIFICACION DE CONDICIONES|\bREV ?\d\b', re.I)
class _M:
    def search(self, t): return _MOD.search(norm(t))
MODIF = _M()
key_nom = lambda s: re.sub(r'\s+', ' ', norm(re.sub(r'\.(gdoc|pdf|docx?|odt)$', '', s, flags=re.I)).replace('/', ' ')).strip()

# 1 · documentos enviados del año
def fecha_nom(nom):
    # ademas de la normal, un año tecleado con un digito de mas ("28 05 20256" en TUBO 5)
    # TUBO 5: "20256" es 2026 con un 5 de mas en medio (la firmada dice 28/05/2026): 202 + ultimo digito
    m = re.search(r'(\d{1,2})[ ./-](\d{1,2})[ ./-](20\d)\d(\d)(?!\d)', nom)
    if m: return f"{m.group(3)}{m.group(4)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    return fecha(nom)
docs = []; todos = {}
for row in csv.reader(open('fuentes.tsv', encoding='utf-8'), delimiter='\t'):
    ruta = row[2][2:] if row[2].startswith('./') else row[2]; nom = ruta.split('/')[-1]
    if nom.lower().endswith(('.ini', '.tmp', '.xlsx', '.gsheet')) or re.search(NO_ES_HOJA, norm(nom)): continue   # viabilidades, pedidos... no son hojas
    f = fecha_nom(nom)
    d = {'ruta': ruta, 'nom': nom, 'fecha': f, 'k': clave(nom), 'et': etiquetas(nom), 'key': key_nom(nom)}
    todos[d['key']] = d     # TODAS las carpetas, para las firmadas cuya enviada se llama con otro año
    if f and f.startswith(A): docs.append(d)

# 2 · el Excel de las automaticas: id de Drive, tipo, bloques, quien
ws = openpyxl.load_workbook('seleccion_encargos.xlsx', read_only=True, data_only=True)['seleccion_encargos']
H = None; xl = []
for r in ws.iter_rows(values_only=True):
    if H is None: H = r; continue
    if isinstance(r[1], datetime.datetime) and r[0]:
        m = re.search(r'id=([\w-]+)', str(r[21] or ''))
        xl.append({'dir': r[0], 'fecha': r[1].strftime('%Y-%m-%d'), 'tipo': r[2], 'id': m.group(1) if m else None,
                   'bloques': [str(H[3 + i]).strip() for i, v in enumerate(r[3:21]) if v is True], 'quien': r[22],
                   'k': clave(str(r[0]))})
# El id de Drive de cada documento, por su TITULO exacto (drive_automaticas.tsv = listado de la
# carpeta en Drive). Antes se adivinaba por fecha y direccion y fallaba cuando ese dia habia
# varias hojas de la misma direccion (proyecto y subvencion, portales...): 375 de 795 en 2026.
import os as _os
id_por_titulo = {}
if _os.path.exists('drive_automaticas.tsv'):
    for r in csv.DictReader(open('drive_automaticas.tsv', encoding='utf-8'), delimiter='\t'):
        if 'document' in (r.get('mimeType') or ''): id_por_titulo.setdefault(key_nom(r['title']), r['id'])
xl_por_id = {x['id']: x for x in xl if x['id']}
def del_excel(d):
    i = id_por_titulo.get(d['key'])
    if i: return xl_por_id.get(i) or {'id': i, 'dir': None, 'tipo': None, 'bloques': [], 'quien': None}
    c = [x for x in xl if x['fecha'] == d['fecha'] and x['k'] and d['k'] and x['k'][1] == d['k'][1] and (x['k'][0] <= d['k'][0] or d['k'][0] <= x['k'][0])]
    if len(c) == 1: return c[0]
    if len(c) > 1:   # varias hojas ese dia en esa direccion: NO se adivina
        return {'id': None, 'dir': c[0]['dir'], 'tipo': None, 'bloques': [], 'quien': None, 'dudoso': len(c)}
    return None

# 3 · un PDF con el mismo nombre que un Google Doc es su exportacion: va con el.
por_key = collections.defaultdict(list)
for d in docs: por_key[d['key']].append(d)
hojas = []
for k, ds in por_key.items():
    principal = next((d for d in ds if not d['nom'].lower().endswith('.pdf')), ds[0])
    h = {'fecha': principal['fecha'], 'k': principal['k'], 'et': principal['et'], 'titulo': principal['nom'],
         'enviados': [d['ruta'] for d in sorted(ds, key=lambda d: d['nom'].lower().endswith('.pdf'))],
         'xl': del_excel(principal), 'firmadas': [], 'datos': None, 'notas': [], 'estado': None}
    hojas.append(h)

# 4 · las firmadas de las tablas (solo las de hoja del año)
TABLAS = ['tanda1', 'tanda2', 'tanda3', '2024', '2023', '2022', 'sfrec', 'sf2024', 'sf2023', 'sf2022']
por_tit = {key_nom(h['titulo']): h for h in hojas}
sueltas = []
for t in TABLAS:
    ws = openpyxl.load_workbook(f'../docs/cruce_hojas_{t}.xlsx')['Cruce']; cols = [c.value for c in ws[1]]
    ix = {c: i for i, c in enumerate(cols)}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[1] not in ('OK', 'ANULADA') or not str(r[ix['Fecha de la hoja (enviada)']] or '').startswith(A): continue
        dato = {'tabla': t, 'n': r[0], 'estado': r[1], 'firmada': 'PRESUPUESTOS FIRMADOS/' + r[ix['Firmada (fichero)']],
                'recibida': r[ix['Fecha recibida (firmada)']], 'total': r[ix['Total base (sin IVA)']], 'conceptos': r[ix['Conceptos']],
                'forma_pago': r[ix['Forma de pago']], 'que': r[ix['Qué se hace']], 'a_quien': r[ix['A quién']],
                'nota': r[ix['Por qué']], 'fecha_hoja': r[ix['Fecha de la hoja (enviada)']], 'dir': r[ix['Dirección (en la hoja)']]}
        datos_de_firmada = 'datos de la propia firmada' in (dato['nota'] or '') or 'datos leídos de la firmada' in (dato['nota'] or '')
        ken = key_nom(str(r[ix['Enviada (documento)']] or ''))
        h = None if datos_de_firmada else por_tit.get(ken)
        if not h and not datos_de_firmada and ken in todos:
            # su enviada existe pero su nombre dice otro año (p. ej. GIRASOL 21 "03 11 2025", REV1 de 2026)
            d0 = todos[ken]
            h = {'fecha': str(dato['fecha_hoja']), 'k': d0['k'], 'et': d0['et'], 'titulo': d0['nom'], 'enviados': [d0['ruta']],
                 'xl': None, 'firmadas': [], 'datos': None, 'notas': [f"el nombre del documento lleva otra fecha ({d0['fecha']}); la hoja es de {dato['fecha_hoja']}"], 'estado': None}
            hojas.append(h); por_tit[ken] = h
        if h: h['firmadas'].append(dato)
        else: sueltas.append(dato)
for s in sueltas:   # hojas que solo conocemos por la firmada
    # antes, buscar su enviada por direccion y fecha de emision (la columna "Enviada" de la tabla
    # puede ser la pareja mala que se corrigio leyendo)
    ks = clave(s['dir'] or s['firmada'].split('/')[-1])
    misma = [h for h in hojas if h['enviados'] and h['fecha'] == str(s['fecha_hoja']) and ks and h['k'] and h['k'][1] == ks[1]
             and (h['k'][0] <= ks[0] or ks[0] <= h['k'][0])]
    if len(misma) == 1:
        misma[0]['firmadas'].append(s); continue
    hojas.append({'fecha': s['fecha_hoja'], 'k': clave(s['dir'] or s['firmada'].split('/')[-1]), 'et': set(), 'titulo': '(sin enviada en la carpeta: datos de la firmada)',
                  'enviados': [], 'xl': None, 'firmadas': [s], 'datos': None, 'notas': [], 'estado': None})

# 5 · versiones: señal explicita en el titulo, misma direccion y concepto, fecha anterior
for h in hojas:
    if not MODIF.search(h['titulo']) or not h['k']: continue
    previas = [g for g in hojas if g is not h and g['k'] and g['k'][1] == h['k'][1] and (g['k'][0] <= h['k'][0] or h['k'][0] <= g['k'][0])
               and (not h['et'] or not g['et'] or g['et'] <= h['et']) and g['fecha'] <= h['fecha'] and not MODIF.search(g['titulo'])]
    if previas:
        p = max(previas, key=lambda g: g['fecha']); h['version_de'] = p
for h in hojas:   # parecido SIN señal: mismo portal y concepto, fechas distintas -> aviso
    if h.get('version_de') or not h['k']: continue
    gem = [g for g in hojas if g is not h and g['k'] and g['k'][1] == h['k'][1] and g['k'][0] == h['k'][0] and g['et'] and g['et'] == h['et'] and g['fecha'] != h['fecha'] and not g.get('version_de')]
    if gem: h['notas'].append('posible versión de/con: ' + '; '.join(sorted({g['titulo'][:60] for g in gem})))

# 5b · lo decidido LEYENDO las "posibles versiones" (pv_resueltas_*.jsonl, instrucciones_versiones.md)
import glob as _glob
por_titulo = {h['titulo']: h for h in hojas}
pv = json.load(open(f'posibles_versiones_{A}.json', encoding='utf-8')) if _glob.glob(f'posibles_versiones_{A}.json') else []
pv_tit = {g['g']: [x['titulo'] for x in g['hojas']] for g in pv}
for f in _glob.glob('pv_resueltas_*.jsonl'):
    for l in open(f, encoding='utf-8'):
        r = json.loads(l); tits = pv_tit.get(r['g'], [])
        for t in tits:
            if t in por_titulo: por_titulo[t]['notas'] = [n for n in por_titulo[t]['notas'] if not n.startswith('posible')]
        if r['decision'] in ('versiones', 'mixto') and r.get('orden'):
            cad = [por_titulo[t] for t in r['orden'] if t in por_titulo]
            for a, b in zip(cad, cad[1:]): b['version_de'] = a
            for h in cad: h['notas'].append('versiones (leído): ' + r['motivo'])
            if r['decision'] == 'mixto':
                for t in tits:
                    if t not in r['orden'] and t in por_titulo: por_titulo[t]['notas'].append('aparte de sus versiones (leído): ' + r['motivo'])
        elif r['decision'] == 'distintas':
            for t in tits:
                if t in por_titulo: por_titulo[t]['notas'].append('hoja distinta, no versión (leído): ' + r['motivo'])
        else:
            for t in tits:
                if t in por_titulo: por_titulo[t]['notas'].append('PARA TI · ¿versión o distinta?: ' + r['motivo'])

# 5c · SUSTITUCIONES (Monica, 7-oct): una hoja (o sus versiones) reemplazada por OTRA(S) hoja(s),
#      p. ej. la combinada de enero por el juego proyecto + subvencion de septiembre.
sus = json.load(open(f'sustituciones_{A}.json', encoding='utf-8'))['sustituciones'] if _glob.glob(f'sustituciones_{A}.json') else []
# + las leidas por los agentes (candidatas_sustitucion_<A>.json -> su_resueltas_*.jsonl)
_cand = {c['g']: c for c in json.load(open(f'candidatas_sustitucion_{A}.json', encoding='utf-8'))} if _glob.glob(f'candidatas_sustitucion_{A}.json') else {}
for f in _glob.glob('su_resueltas_*.jsonl'):
    for l in open(f, encoding='utf-8'):
        r = json.loads(l); c = _cand.get(r['g'])
        if not c: continue
        if r['decision'] == 'sustituida' and r.get('por'):
            sus.append({'antiguas': [c['antigua']['titulo']], 'nuevas': r['por'], 'motivo': '(leído) ' + r['motivo']})
        elif r['decision'] == 'no se sabe' and c['antigua']['titulo'] in por_titulo:
            por_titulo[c['antigua']['titulo']]['notas'].append('PARA TI · ¿sustituida por una posterior?: ' + r['motivo'])
for sx in sus:
    nuevas = [por_titulo[t] for t in sx['nuevas'] if t in por_titulo]
    for t in sx['antiguas'] + sx['nuevas']:
        h = por_titulo.get(t)
        if h:
            h['notas'] = [n for n in h['notas'] if 'versión' not in n and 'versiones' not in n and not n.startswith('PARA TI')]
            if h in nuevas and h.get('version_de') and h['version_de']['titulo'] in sx['antiguas']: h.pop('version_de')
    for t in sx['antiguas']:
        h = por_titulo.get(t)
        if h: h['sustituida_por'] = nuevas; h['notas'].append('SUSTITUIDA: ' + sx['motivo'])
    for h in nuevas: h['notas'].append('sustituye a la(s) hoja(s) anterior(es): ' + sx['motivo'])

# 6 · cadenas de versiones -> una HOJA con v1..vn
raiz = lambda h: raiz(h['version_de']) if h.get('version_de') else h
grupos = collections.defaultdict(list)
for h in hojas: grupos[id(raiz(h))].append(h)
filas = []
for g in grupos.values():
    g.sort(key=lambda h: (h['fecha'], h['titulo']))
    filas.append(g)
filas.sort(key=lambda g: (g[0]['fecha'], g[0]['titulo']))
# cada firmada, a la version con SU fecha de emision (la de la tabla), si la hay
for g in filas:
    if len(g) < 2: continue
    firm = [d for h in g for d in h['firmadas']]
    for h in g: h['firmadas'] = []
    for d in firm:
        mismas = [h for h in g if h['fecha'] == str(d['fecha_hoja'])]
        (mismas[0] if mismas else g[0])['firmadas'].append(d)   # misma fecha: el original (leido en Castellon 15 e Isaac Peral 5)

# 7 · el Excel
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Hojas ' + A
cols = ['Código provisional', 'Estado', 'Fecha (v1)', 'Dirección', 'Tipo (Excel automáticas)', 'Bloques marcados (Excel)',
        'Versiones', 'Ficheros ENVIADOS', 'Ficheros FIRMADOS', 'Fecha recibida', 'Total base (firmada)', 'Conceptos (firmada)',
        'Forma de pago', 'A quién', 'Generada por (solo cotejo)', 'Id Drive (última versión)', 'Notas', 'Opp', 'TU DECISIÓN']
ws.append(cols); cuenta = collections.Counter()
codigo = {id(h): f'HE-{A}-{i:04d}' for i, g in enumerate(filas, 1) for h in g}
for i, g in enumerate(filas, 1):
    ult = g[-1]; firm = [d for h in g for d in h['firmadas']]
    anulada = any(d['estado'] == 'ANULADA' for d in firm)
    estado = 'ANULADA' if anulada else 'FIRMADA' if firm else 'ENVIADA SIN FIRMAR'
    sp = next((h['sustituida_por'] for h in g if h.get('sustituida_por')), None)
    if sp and not firm:
        estado = 'SUSTITUIDA'
        g[-1]['notas'].append('sustituida por: ' + ', '.join(sorted({codigo[id(n)] for n in sp})))
    cuenta[estado] += 1
    versiones = '\n'.join(f"v{j} {h['fecha']} · {h['titulo'][:70]}" + ('  ← FIRMADA' if h['firmadas'] else '') for j, h in enumerate(g, 1)) if len(g) > 1 else ''
    enviados = '\n'.join((f"v{j}: " if len(g) > 1 else '') + e for j, h in enumerate(g, 1) for e in h['enviados'])
    firmados = '\n'.join(d['firmada'] for d in firm)
    x = ult['xl'] or next((h['xl'] for h in reversed(g) if h['xl']), None)
    d0 = firm[0] if firm else {}
    notas = [n for h in g for n in h['notas']] + [f"{d['tabla']} {d['n']}: {d['nota']}" for d in firm if d.get('nota')]
    if x and x.get('dudoso'):
        notas.append(f"id de Drive POR CONFIRMAR: ese día hay {x['dudoso']} hojas de esta dirección en el Excel y el título no casó")
    ws.append([f'HE-{A}-{i:04d}', estado, g[0]['fecha'], (d0.get('dir') or (x or {}).get('dir') or ''),
               (x or {}).get('tipo'), ', '.join((x or {}).get('bloques') or []), versiones, enviados, firmados,
               ', '.join(sorted({str(d['recibida']) for d in firm if d['recibida']})), d0.get('total'), d0.get('conceptos'),
               d0.get('forma_pago'), d0.get('a_quien'), (x or {}).get('quien'), (x or {}).get('id'), '\n'.join(notas)[:1500], '', ''])
color = {'FIRMADA': 'DCEFC8', 'ENVIADA SIN FIRMAR': 'FFF2CC', 'ANULADA': 'D9D9D9', 'SUSTITUIDA': 'E8E0F0'}
for j, a in enumerate([16, 18, 11, 30, 26, 34, 50, 60, 55, 12, 12, 40, 40, 18, 12, 22, 60, 10, 30], 1):
    ws.column_dimensions[get_column_letter(j)].width = a
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='2B2B2B'); c.alignment = Alignment(wrap_text=True)
ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
for f in ws.iter_rows(min_row=2):
    for c in f: c.alignment = Alignment(wrap_text=True, vertical='top')
    f[1].fill = PatternFill('solid', fgColor=color[f[1].value]); f[18].fill = PatternFill('solid', fgColor='EAF4FF')
wb.save(f'../docs/manifiesto_hojas_{A}.xlsx')
print(A, 'hojas', len(filas), dict(cuenta), '| con varias versiones', sum(1 for g in filas if len(g) > 1),
      '| posibles versiones sin señal', sum(1 for g in filas for h in g if h['notas'] and h['notas'][0].startswith('posible')),
      '| firmadas sueltas (sin enviada)', len(sueltas), '| sin Excel', sum(1 for g in filas if not any(h['xl'] for h in g)))

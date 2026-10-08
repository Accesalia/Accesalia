# REVISION DE LAS FIRMADAS: las tablas de trabajo a Excel para que Monica lo vea (solo LEE).
#   python exportar_revision.py  -> docs/revision_firmadas.xlsx (Firmadas, Lineas, Plazos de cobro)
import sys, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
sys.path.insert(0, '../scripts')
import produccion
base = produccion.arrancar() or produccion.base

F = base.leer('revision_firmadas?select=*,opp:oportunidad_id(codigo,nombre,comunidad_provisional),'
              'contrata:contrata_id(nombre)&order=numero_hoja')
L = base.leer('revision_firmadas_lineas?select=*,bloque:bloque_id(codigo)&order=orden,id')
P = base.leer('revision_firmadas_plazos?select=*&order=orden,id')
num = {f['id']: f['numero_hoja'] for f in F}
lin = {l['id']: l for l in L}
COLOR = {'ok': 'D9EAD3', 'pendiente': 'FFF2CC', 'duda': 'F4CCCC', 'corregido': 'CFE2F3',
         'coincide': 'D9EAD3', 'distinta': 'FFF2CC', 'falta_en_base': 'F4CCCC', 'sobra_en_base': 'F4CCCC'}
HITO = {'firma': 'A la firma', 'encargo': 'Al encargo', 'entrega': 'A la entrega', 'licencia': 'A la licencia',
        'cfo': 'Al fin de obra', 'concesion': 'A la concesion', 'otro': 'Otro'}

wb = openpyxl.Workbook()
def hoja(ws, cab, filas, anchos, col_color):
    ws.append(cab)
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='3C3C3C')
    for f in filas:
        ws.append(f)
        k = f[col_color] if col_color < len(f) else None
        if k in COLOR: ws.cell(ws.max_row, col_color + 1).fill = PatternFill('solid', fgColor=COLOR[k])
    for i, a in enumerate(anchos, 1): ws.column_dimensions[get_column_letter(i)].width = a
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'B2'

ws = wb.active; ws.title = 'Firmadas'
hoja(ws, ['Hoja', 'Revision', 'Opp', 'Direccion (opp)', 'Emision', 'Firma', 'Dias hasta firmar', 'Fecha en la casilla',
          'Firmada', 'Paga', 'Razon social', 'CIF', 'IBAN', 'Ya existe en la app', 'Total sin IVA', 'Forma de pago (literal)', 'Notas'],
     [[f['numero_hoja'], f['revision'], f['opp']['codigo'], f['opp']['nombre'] or f['opp']['comunidad_provisional'],
       f['fecha_emision'], f['fecha_firma'],
       (__import__('datetime').date.fromisoformat(f['fecha_firma']) - __import__('datetime').date.fromisoformat(f['fecha_emision'])).days
       if f['fecha_firma'] and f['fecha_emision'] else None,
       f['fecha_casilla'], 'si' if f['firma_presente'] else 'NO', f['pagador_tipo'], f['pagador_razon_social'], f['pagador_cif'],
       f['pagador_iban'], ('si: ' + f['contrata']['nombre']) if f['contrata'] else 'si' if f['comunidad_id'] else
       ('NO, hay que crearla' if f['pagador_hay_que_crear'] else '-'),
       f['total_base'], f['forma_pago_texto'], f['nota']] for f in F],
     [13, 10, 13, 30, 11, 11, 8, 14, 8, 10, 26, 13, 28, 16, 10, 50, 70], 1)
hoja(wb.create_sheet('Lineas'), ['Hoja', 'Comparacion con la base', 'Orden', 'Concepto (texto de la hoja)', 'Bloque', 'Importe fijo',
                                 '% a exito', 'Incluida (sin precio propio)', 'Nota'],
     [[num[l['firmada_id']], l['comparacion'], l['orden'], l['texto'], (l['bloque'] or {}).get('codigo'), l['importe'],
       l['porcentaje'], 'si' if l['incluido'] else '', l['nota']] for l in sorted(L, key=lambda l: (num[l['firmada_id']], l['orden'] or 99))],
     [13, 14, 6, 45, 26, 11, 9, 10, 40], 1)
hoja(wb.create_sheet('Plazos de cobro'), ['Hoja', 'Concepto', 'Importe de la linea', 'Cuando', 'Plazo', '% de la linea', 'Importe', 'Texto de la hoja'],
     [[num[lin[p['linea_id']]['firmada_id']], lin[p['linea_id']]['texto'], lin[p['linea_id']]['importe'], HITO[p['hito']], p['orden'],
       p['porcentaje'], p['importe'], p['texto']]
      for p in sorted(P, key=lambda p: (num[lin[p['linea_id']]['firmada_id']], lin[p['linea_id']]['orden'] or 99, p['orden']))],
     [13, 40, 11, 14, 6, 9, 10, 60], 99)
SALIDA = sys.argv[1] if len(sys.argv) > 1 else '../docs/revision_firmadas.xlsx'
wb.save(SALIDA)
print(SALIDA + ':', len(F), 'firmadas,', len(L), 'lineas,', len(P), 'plazos')
